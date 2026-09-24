"""Uji layanan: nilai ulang saat submit, sinkron Google Sheet (outbox + retry + batch), kunci, admin, ekspor, pembersihan.

Google Sheet dipalsukan (FakeSheets) — tidak ada panggilan jaringan.
"""
import csv
import datetime as dt
import io
import os
import time
import unittest
import zipfile

from fastapi.testclient import TestClient

from server import config, services, sheets
from server.db import ScanSession, SessionLocal
from server.main import app
from tests.regression.fixtures import KUNCI

ROOT = os.path.join(os.path.dirname(__file__), "..", "..")
ADM = {"X-Admin-Token": "rahasia"}
VALID = {"nama_pengawas": "Budi Santoso", "hp": "081234567890", "ruangan": "TULT 0603", "kelas": "BS1SI-50-REG-01",
         "prodi": "S1 Sistem Informasi"}
PDF = open(os.path.join(ROOT, "LJK.pdf"), "rb").read()


class FakeSheets:
    def __init__(self):
        self.calls, self.fail, self.kunci = [], 0, {}

    def append_records(self, recs):
        if self.fail > 0:
            self.fail -= 1
            raise RuntimeError("429 quota exceeded")
        self.calls.append(list(recs))
        return len(recs)

    def fetch_kunci(self):
        return self.kunci


class TestServices(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from server import refdata
        cls._orig_kp, refdata.kelas_prodi = refdata.kelas_prodi, lambda k: VALID["prodi"]   # tes memakai nama kelas bebas
        cls._ctx = TestClient(app)
        cls.c = cls._ctx.__enter__()

    @classmethod
    def tearDownClass(cls):
        from server import refdata
        refdata.kelas_prodi = cls._orig_kp
        sheets.set_client(None)
        cls._ctx.__exit__(None, None, None)

    def setUp(self):
        self.fake = FakeSheets()
        sheets.set_client(self.fake)
        self.drain()

    def drain(self):
        """Isolasi antar tes: anggap semua sesi lama sudah tersinkron & hapus semua kunci."""
        from server.db import Kunci
        with SessionLocal() as db:
            for s in db.query(ScanSession).filter(ScanSession.submitted.is_(True), ScanSession.synced_at.is_(None)):
                s.synced_at = dt.datetime.now(dt.timezone.utc)
            db.query(Kunci).delete()
            db.commit()
        self.fake.calls.clear()

    def submitted_session(self, kelas="IF-47-01", n=1):
        sid = self.c.post("/api/sessions", json={**VALID, "kelas": kelas}).json()["id"]
        self.c.post(f"/api/sessions/{sid}/files", files=[("files", (f"l{i}.pdf", PDF, "application/pdf")) for i in range(n)])
        t = time.time()
        while time.time() - t < 120 and self.c.get(f"/api/sessions/{sid}").json()["scanning"]:
            time.sleep(0.4)
        self.c.post(f"/api/sessions/{sid}/validate-all")
        return sid

    def submit(self, sid):
        r = self.c.post(f"/api/sessions/{sid}/submit")
        self.assertEqual(r.status_code, 200, r.text)

    def test_admin_requires_token(self):
        for path in ("/api/admin/summary", "/api/admin/sessions", "/api/admin/kunci", "/api/admin/export.xlsx"):
            self.assertEqual(self.c.get(path).status_code, 401, path)
            self.assertEqual(self.c.get(path, headers={"X-Admin-Token": "salah"}).status_code, 401, path)

    def test_regrade_at_submit_uses_latest_key(self):
        sid = self.submitted_session()   # dipindai tanpa kunci
        before = self.c.get("/api/admin/export.csv", headers=ADM).text
        self.c.put("/api/admin/kunci/A", json={str(q): a for q, a in KUNCI["A"].items()}, headers=ADM)   # kunci menyusul
        self.submit(sid)
        rows = [r for r in csv.DictReader(io.StringIO(self.c.get(f"/api/admin/export.csv?sid={sid}", headers=ADM).text))]
        self.assertEqual(len(rows), 1)
        self.assertNotIn(rows[0]["Nilai"], ("-", ""), rows[0])
        self.assertEqual(rows[0]["Kunci Terpakai"], "A")

    def test_sync_batches_all_due_sessions_in_one_call(self):
        s1, s2 = self.submitted_session("K1"), self.submitted_session("K2", n=2)
        self.submit(s1); self.submit(s2)
        r = services.sync_pending_once()
        self.assertEqual((r["synced"], len(self.fake.calls)), (2, 1))          # SATU permintaan untuk 2 sesi
        self.assertEqual(len(self.fake.calls[0]), 3)                           # 1 + 2 lembar
        self.assertEqual({x["Kelas"] for x in self.fake.calls[0]}, {"K1", "K2"})
        self.assertEqual(services.sync_pending_once()["synced"], 0)            # tidak dikirim dua kali
        items = self.c.get("/api/admin/sessions?status=submitted", headers=ADM).json()["items"]
        self.assertTrue(all(i["synced"] for i in items if i["id"] in (s1, s2)))

    def test_sync_failure_backoff_then_recover(self):
        sid = self.submitted_session("K3")
        self.submit(sid)
        self.fake.fail = 1
        r = services.sync_pending_once()
        self.assertEqual((r["synced"], r["failed"]), (0, 1))
        me = next(i for i in self.c.get("/api/admin/sessions?status=unsynced", headers=ADM).json()["items"] if i["id"] == sid)
        self.assertEqual((me["synced"], me["sync_attempts"]), (False, 1))
        self.assertIn("429", me["sync_error"])
        self.assertEqual(services.sync_pending_once()["synced"], 0)            # masih masa tunggu (backoff)
        self.assertEqual(self.c.post("/api/admin/sync/run", headers=ADM).json()["synced"], 1)   # admin paksa ulang
        self.assertEqual(len(self.fake.calls), 1)

    def test_poison_session_does_not_block_others(self):
        good, bad = self.submitted_session("OK1"), self.submitted_session("BAD1")
        self.submit(good); self.submit(bad)
        orig = self.fake.append_records

        def picky(recs):
            if any(r.get("Kelas") == "BAD1" for r in recs):
                raise RuntimeError("baris tidak valid")
            return orig(recs)
        self.fake.append_records = picky
        r = services.sync_pending_once()
        self.assertEqual((r["synced"], r["failed"]), (1, 1))
        st = {i["id"]: i["synced"] for i in self.c.get("/api/admin/sessions", headers=ADM).json()["items"]}
        self.assertEqual((st[good], st[bad]), (True, False))

    def test_sync_since_skips_older_sessions(self):
        sid = self.submitted_session("SINCE1")
        self.submit(sid)
        old = config.SYNC_SINCE
        try:
            config.SYNC_SINCE = (dt.datetime.now(dt.timezone.utc) + dt.timedelta(hours=1)).isoformat()
            self.assertEqual(services.sync_pending_once()["synced"], 0)        # sesi lebih lama dari batas -> dilewati
            self.assertEqual(len(self.fake.calls), 0)
            config.SYNC_SINCE = (dt.datetime.now(dt.timezone.utc) - dt.timedelta(hours=1)).isoformat().replace("+00:00", "Z")
            self.assertEqual(services.sync_pending_once()["synced"], 1)        # setelah batas -> terkirim
        finally:
            config.SYNC_SINCE = old

    def test_toml_credentials_reuse_streamlit_secrets(self):
        import tempfile
        with tempfile.NamedTemporaryFile("w", suffix=".toml", delete=False) as f:
            f.write('[connections.gsheets]\nspreadsheet = "https://docs.google.com/spreadsheets/d/ABC/edit"\n'
                    'type = "service_account"\nproject_id = "p"\nclient_email = "x@p.iam.gserviceaccount.com"\nprivate_key = "K"\n')
        creds, url = sheets.load_toml_credentials(f.name)
        os.unlink(f.name)
        self.assertEqual(url, "https://docs.google.com/spreadsheets/d/ABC/edit")
        self.assertEqual(creds["client_email"], "x@p.iam.gserviceaccount.com")
        self.assertNotIn("spreadsheet", creds)                                  # bukan bagian kredensial

    def test_sync_not_configured_is_noop(self):
        sheets.set_client(None)
        self.assertEqual(services.sync_pending_once()["configured"], False)
        self.assertFalse(self.c.get("/api/admin/summary", headers=ADM).json()["gsheet_configured"])

    def test_kunci_sync_upload_delete(self):
        self.fake.kunci = {"kj048": {1: "A", 2: "B", 3: "C", 4: "D"}}
        r = self.c.post("/api/admin/kunci/sync", headers=ADM).json()
        self.assertEqual(r["kunci"], 1)
        lst = {k["name"]: k for k in self.c.get("/api/admin/kunci", headers=ADM).json()}
        self.assertEqual((lst["kj048"]["soal"], lst["kj048"]["source"]), (4, "gsheet"))
        up = self.c.post("/api/admin/kunci/upload", headers=ADM, files={"file": ("kj999.csv", b"1,A\n2,B\n3,C\n4,D\n5,A\n", "text/csv")})
        self.assertEqual(up.status_code, 200, up.text)
        self.assertEqual(up.json()["kunci"], {"kj999": 5})
        self.assertEqual(self.c.post("/api/admin/kunci/upload", headers=ADM, files={"file": ("x.csv", b"halo", "text/csv")}).status_code, 422)
        self.assertEqual(self.c.post("/api/admin/kunci/upload", headers=ADM, files={"file": ("x.txt", b"1,A", "text/plain")}).status_code, 415)
        self.assertEqual(self.c.delete("/api/admin/kunci/kj999", headers=ADM).status_code, 204)
        self.assertEqual(self.c.delete("/api/admin/kunci/kj999", headers=ADM).status_code, 404)

    def test_export_xlsx_and_filter(self):
        sid = self.submitted_session("XL-01")
        self.submit(sid)
        r = self.c.get("/api/admin/export.xlsx?kelas=XL-01", headers=ADM)
        self.assertEqual(r.status_code, 200)
        self.assertTrue(zipfile.is_zipfile(io.BytesIO(r.content)))
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(r.content)).active
        header = [c.value for c in ws[1]]
        self.assertEqual(header[:5], ["Submit Date", "Nama Pengawas", "No HP Pengawas", "Ruangan", "Kelas"])
        self.assertEqual({row[header.index("Kelas")].value for row in ws.iter_rows(min_row=2)}, {"XL-01"})

    def test_regrade_all_after_key_change(self):
        sid = self.submitted_session("RG-01")
        self.submit(sid)                                   # disubmit TANPA kunci -> Nilai "-"
        def nilai():
            rows = list(csv.DictReader(io.StringIO(self.c.get(f"/api/admin/export.csv?sid={sid}", headers=ADM).text)))
            return rows[0]["Nilai"], rows[0]["Kunci Terpakai"]
        self.assertEqual(self.c.post("/api/admin/regrade", headers=ADM).json()["error"], "Belum ada kunci jawaban")
        self.assertEqual(nilai()[0], "-")
        self.c.put("/api/admin/kunci/A", json={str(q): a for q, a in KUNCI["A"].items()}, headers=ADM)
        r1 = self.c.post("/api/admin/regrade?kelas=RG-01", headers=ADM).json()
        self.assertEqual((r1["sessions"], r1["changed"], r1["changed_in_synced"]), (1, 1, 0))
        self.assertNotEqual(nilai()[0], "-")
        self.assertEqual(self.c.post("/api/admin/regrade?kelas=RG-01", headers=ADM).json()["changed"], 0)   # idempoten
        # kunci berubah -> nilai berubah; sesi yang sudah terkirim ke Sheet dilaporkan usang
        services.sync_pending_once()
        before = nilai()[0]
        self.c.put("/api/admin/kunci/A", json={str(q): "D" for q in range(1, 76)}, headers=ADM)
        r2 = self.c.post("/api/admin/regrade?kelas=RG-01", headers=ADM).json()
        self.assertEqual((r2["changed"], r2["changed_in_synced"]), (1, 1))
        self.assertNotEqual(nilai()[0], before)
        # filter kelas lain tidak tersentuh
        self.assertEqual(self.c.post("/api/admin/regrade?kelas=TIDAK-ADA", headers=ADM).json()["sessions"], 0)

    def test_regrade_pull_from_sheet(self):
        self.fake.kunci = {"kj777": {q: "A" for q in range(1, 11)}}
        r = self.c.post("/api/admin/regrade?pull=true", headers=ADM).json()
        self.assertEqual((r["gsheet_configured"], r["kunci_ditarik"]), (True, 1))

    def test_record_key_order_survives_database_roundtrip(self):
        """Urutan kolom rekap (ekspor/Sheet) bergantung pada urutan kunci JSON. PostgreSQL `jsonb` mengacaknya; kolom JSON biasa tidak."""
        from server.db import Sheet, UploadFile, engine
        sid = self.submitted_session("ORD-1")
        order = ["Zebra", "Alpha", "Nama Pengawas", "NPM", "b", "aa", "Kelas"]      # sengaja bukan alfabet/panjang
        with SessionLocal() as db:
            sh = db.query(Sheet).filter(Sheet.session_id == sid).first()
            sh.record = {k: "x" for k in order}
            db.commit()
        with SessionLocal() as db:
            got = list(db.query(Sheet).filter(Sheet.session_id == sid).first().record.keys())
        self.assertEqual(got, order, f"urutan kunci berubah di {engine.dialect.name}")

    def test_cleanup_removes_old_upload_folders_only(self):
        old, fresh = self.submitted_session("OLD"), self.submitted_session("NEW")
        self.submit(old); self.submit(fresh)
        with SessionLocal() as db:
            db.get(ScanSession, old).submitted_at = dt.datetime.now(dt.timezone.utc) - dt.timedelta(hours=config.UPLOAD_RETENTION_HOURS + 2)
            db.commit()
        self.assertTrue(os.path.isdir(os.path.join(config.UPLOAD_DIR, old)))
        services.cleanup_once()
        self.assertFalse(os.path.isdir(os.path.join(config.UPLOAD_DIR, old)))
        self.assertTrue(os.path.isdir(os.path.join(config.UPLOAD_DIR, fresh)))


if __name__ == "__main__":
    unittest.main()
